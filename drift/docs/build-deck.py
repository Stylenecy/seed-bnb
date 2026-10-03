"""DRIFT pitch deck builder (python-pptx).

Uses only verified facts: the real Dex deployment, its receipts (docs/deployment-dex.md),
test counts from real runs, and the provenance sentence read from README.md.
Business-model statements are labelled as hypotheses; no market figures are used.
Output: docs/submission/DRIFT-pitch.pptx (a PDF with embedded fonts is exported from it).

Visual rules (docs/VISUAL-DIRECTION.md): colours are the brand tokens; gold marks only what
a judge can check on-chain (addresses, block numbers, the halt line, "exact match", the
BscScan and Sourcify links); periwinkle marks only the off-chain engine. Headings are white.
Fonts: Calibri and Consolas stand in for Geist and Geist Mono, because a .pptx renders
with the fonts installed on the viewer's machine and Geist is not a system font. Mono is
used for machine-written values: in box() text, a `backticked` span becomes a mono run.
"""
import os
import re
import zipfile
from datetime import datetime, timezone

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))  # drift/docs
DRIFT = os.path.dirname(HERE)
OUT = os.path.join(DRIFT, "docs", "submission", "DRIFT-pitch.pptx")
SHOTS = os.path.join(HERE, "screens", "after")

# Same wording as README.md "This fork's contribution" (read from it, so the two never drift apart).
with open(os.path.join(DRIFT, "README.md"), encoding="utf-8") as fh:
    PROVENANCE = re.search(r"^DRIFT's core \(the quant engine.*$", fh.read(), re.M).group(0)

CONTRACT = "0x8b09ebB85Be8Ed55Bb5132d29eABc567c42aa83D"
LIVE_URL = "https://drift-macroguard.vercel.app/macroguard"
REPO_URL = "https://github.com/Stylenecy/seed-bnb/tree/dex/drift/drift"
BSCSCAN_URL = f"https://testnet.bscscan.com/address/{CONTRACT}"
SOURCIFY_URL = f"https://repo.sourcify.dev/97/{CONTRACT}"

SANS = "Calibri"
MONO = "Consolas"


def hex_rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


def over_ink(value: str, alpha: float) -> RGBColor:
    """A token at `alpha` opacity on the ink ground, as the web composes it (e.g. white/65)."""
    top = [int(value.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4)]
    ink = [0x0B, 0x0C, 0x0F]
    return RGBColor(*[round(t * alpha + b * (1 - alpha)) for t, b in zip(top, ink)])


# Brand tokens (VISUAL-DIRECTION.md)
BG = hex_rgb("#0B0C0F")  # ink
FG = hex_rgb("#FFFFFF")
MUTED = over_ink("#FFFFFF", 0.65)  # body text >= white/60
PANEL = over_ink("#FFFFFF", 0.06)
GOLD = hex_rgb("#F0B90B")  # on-chain / verifiable
ACCENT = hex_rgb("#9AA8F0")  # periwinkle: off-chain engine
GREEN = hex_rgb("#34D399")  # allowed / running / passing
AMBER = hex_rgb("#FBBF24")  # attention
ROSE = hex_rgb("#FB7185")  # halted / blocked

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def slide(notes=None):
    s = prs.slides.add_slide(blank)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def box(s, left, top, width, height, text, size=24, color=FG, bold=False, font=SANS):
    """Text box; each line is a paragraph, and each `backticked` span is set in mono."""
    tx = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tx.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.35)
        for part in re.split(r"(`[^`]*`)", line):
            if not part:
                continue
            mono = len(part) >= 2 and part[0] == part[-1] == "`"
            run = p.add_run()
            run.text = part[1:-1] if mono else part
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = MONO if mono else font
    return tx


def link_row(s, top, label, url, color, clickable=False):
    """A muted label and a full URL in mono; on-chain links are clickable (gold, see the theme fix)."""
    box(s, 0.7, top + 0.03, 1.9, 0.45, label, 13, MUTED)
    tx = s.shapes.add_textbox(Inches(2.5), Inches(top), Inches(10.3), Inches(0.45))
    run = tx.text_frame.paragraphs[0].add_run()
    run.text = url
    if clickable:
        run.hyperlink.address = url
    run.font.size = Pt(14)
    run.font.name = MONO
    run.font.color.rgb = color


def rect(s, left, top, width, height, color, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    return r


def shot(s, name, left, top, width, crop=None):
    """Picture at `width`; `crop` = (left, top, right, bottom) fractions to trim to the panel."""
    path = os.path.join(SHOTS, name)
    if not os.path.exists(path):
        raise SystemExit(f"missing screenshot: {path}")
    pic = s.shapes.add_picture(path, Inches(left), Inches(top))
    if crop:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        visible_w = pic.width * (1 - crop[0] - crop[2])
        visible_h = pic.height * (1 - crop[1] - crop[3])
    else:
        visible_w, visible_h = pic.width, pic.height
    pic.width = Inches(width)
    pic.height = int(Inches(width) * visible_h / visible_w)
    return pic


def title(s, text):
    box(s, 0.7, 0.4, 11.9, 0.9, text, 36, FG, True)


# 1. Title
s = slide("DRIFT in one line: the trading bot's risk rules live in a public contract on BNB Chain, so anyone can check them.")
box(s, 0.7, 0.7, 11.9, 1.2, "DRIFT", 60, FG, True)
box(s, 0.7, 2.0, 11.9, 1.2, "A trading bot whose risk rules you can verify", 34, FG, True)
box(s, 0.7, 3.2, 11.9, 0.8, "Off-chain quant research · on-chain risk gate on BNB Chain", 22, MUTED)
box(s, 0.7, 4.6, 11.9, 0.5, "Live demo", 16, MUTED)
box(s, 0.7, 5.0, 11.9, 0.5, LIVE_URL.removeprefix("https://"), 18, FG, font=MONO)
box(s, 0.7, 5.6, 11.9, 0.5, "Dex's contract on BSC Testnet", 16, MUTED)
box(s, 0.7, 6.0, 11.9, 0.5, CONTRACT, 18, GOLD, font=MONO)
box(s, 0.7, 6.75, 11.9, 0.5, "Indonesia Web3 Hackathon 2026 · Dex Bennett", 16, MUTED)

# 2. Problem and promise
s = slide("Bots ask for trust. DRIFT replaces the promise with a public rule and a public record.")
title(s, "The problem")
box(s, 0.7, 1.6, 11.9, 1.6, "Trading bots ask users to trust risk rules nobody can see.\nSettings change silently; off-chain decision logs can be rewritten.", 24)
box(s, 0.7, 3.6, 11.9, 0.7, "The promise", 28, FG, True)
box(s, 0.7, 4.4, 11.9, 1.6, "Anyone can check what DRIFT's bot is allowed to do right now,\nand every recorded decision is public on BscScan.", 24)

# 3. Who needs it and the business model (hypotheses)
s = slide(
    "Everything on this slide is a hypothesis: no user interviews or market sizing have been done yet. "
    "The point is the shape of the value: proof of discipline, not proof of profit."
)
title(s, "Who needs it, and how it could pay")
box(s, 0.7, 1.35, 11.9, 0.5, "HYPOTHESIS: not yet validated with users; no market figures claimed", 15, AMBER, True)
box(s, 0.7, 2.0, 5.6, 0.6, "Who needs it", 24, FG, True)
box(
    s, 0.7, 2.7, 5.6, 3.6,
    "Retail traders who follow or rent bots: proof that the risk stop is real\n"
    "Bot builders and signal sellers: a record of discipline a screenshot cannot fake\n"
    "Communities and auditors: check a bot's rules without trusting its operator",
    18,
)
box(s, 6.9, 2.0, 5.7, 0.6, "How it could pay", 24, FG, True)
box(
    s, 6.9, 2.7, 5.7, 3.6,
    "A risk gate per bot: its own MacroGuard plus a public panel, as a monthly service\n"
    "The panel link as a trust badge for strategy sellers\n"
    "No custody and no performance fee; on-chain cost per decision stays tiny (slide 7)",
    18,
)

# 4. How it works
s = slide("Four steps: signal off-chain, ask the gate, trade on the exchange, record on-chain.")
title(s, "How DRIFT works")
steps = [
    ("1  Signal", "Python engine: four strategies, point-in-time, no look-ahead", ACCENT),
    ("2  Ask the gate", "`MacroGuard.allowed(signal)` on BSC Testnet: regime + 20% halt", FG),
    ("3  Trade", "Runner places the order on Bybit testnet (vetoed signals go Flat)", ACCENT),
    ("4  Record", "`recordDecision(...)` leaves a public receipt; a breach halts it", FG),
]
for i, (head, body, color) in enumerate(steps):
    y = 1.6 + i * 1.15
    box(s, 0.7, y, 3.2, 0.6, head, 24, color, True)
    box(s, 4.0, y + 0.05, 8.6, 0.9, body, 20, FG)
box(s, 0.7, 6.1, 11.9, 0.5, "Periwinkle: off-chain engine · white: on-chain contract", 14, MUTED)
box(s, 0.7, 6.5, 11.9, 0.6, "The contract never executes orders; it is the public rule and the public record.", 18, MUTED)

# 5. Product
s = slide(
    "Public, no login, no wallet. Left: the live state read from BNB Chain in the browser. "
    "Right: ask the contract a what-if; it answers through eth_call, nothing is signed."
)
title(s, "The panel, and asking the contract")
# Trim the site navigation and the page margins so only the panel shows.
shot(s, "public-guard-desktop.png", 0.5, 1.4, 6.1, crop=(0.097, 0.094, 0.097, 0.12))
shot(s, "ask-contract-desktop.png", 6.8, 1.4, 6.0)
box(s, 0.5, 5.25, 6.1, 1.2, "Live regime, halt, 20% line and decisions,\nread from BSC Testnet at a named block", 16, MUTED)
box(s, 6.8, 5.25, 6.0, 1.2, "Long at −25%: Blocked, answered by the live contract\nvia `eth_call` from the agent address (simulation)", 16, MUTED)

# 6. The halt story
s = slide(
    "Smoke test on 30 Sep 2026. At step 3 nobody pressed a button: the rule fired inside the contract "
    "when the reported drawdown crossed 20%. Honest limit: the drawdown was reported by the agent."
)
title(s, "−25% recorded: the contract halted itself")
# gauge: 0 → 30%, halt line at 20%, recorded −1% and −25%
gx, gy, gw = 0.9, 2.05, 11.5
rect(s, gx, gy, gw * 20 / 30, 0.22, over_ink("#34D399", 0.25))
rect(s, gx + gw * 20 / 30, gy, gw * 10 / 30, 0.22, over_ink("#FB7185", 0.25))
rect(s, gx + gw * 20 / 30 - 0.02, gy - 0.25, 0.05, 0.72, GOLD)
for pct, color in ((1, GREEN), (25, ROSE)):
    rect(s, gx + gw * pct / 30 - 0.17, gy - 0.06, 0.34, 0.34, color, MSO_SHAPE.OVAL)
box(s, gx + gw * 20 / 30 - 0.9, gy - 0.75, 2.2, 0.4, "halt line 20%", 14, GOLD, True)
box(s, gx - 0.1, gy + 0.35, 3.0, 0.4, "−1%: kept running", 14, GREEN)
box(s, gx + gw * 25 / 30 - 1.6, gy + 0.35, 3.4, 0.4, "−25%: halted itself", 14, ROSE)
rows = [
    ("134,042,196", "setRegime(RiskOff)", "Long blocked; Short and Flat open", MUTED),
    ("134,042,256", "recordDecision(BNB, Short, −1%)", "Decision #1, inside the line", GREEN),
    ("134,042,283", "recordDecision(BNB, Short, −25%)", "Halted event: only Flat allowed", ROSE),
    ("134,042,318", "resume()", "Agent-only, on the record", MUTED),
    ("134,042,328", "setRegime(Neutral)", "Long allowed again; 2 decisions on-chain", MUTED),
]
for i, (block, call, outcome, color) in enumerate(rows):
    y = 3.15 + i * 0.68
    if color == ROSE:
        rect(s, 0.6, y - 0.05, 12.1, 0.64, PANEL)
    box(s, 0.8, y, 2.3, 0.5, f"#{block}", 16, GOLD, font=MONO)
    box(s, 3.2, y, 4.6, 0.5, call, 16, FG, color == ROSE, font=MONO)
    box(s, 7.9, y, 4.8, 0.5, outcome, 17, color)
box(s, 0.7, 6.75, 11.9, 0.5, "Every row is a transaction with receipt status 1 on BscScan · block `134,042,283` is the halt", 14, MUTED)

# 7. Why BNB Chain
s = slide(
    "Costs come from our own receipts on BSC Testnet at the observed 0.1 gwei; they are not a mainnet quote. "
    "The panel reads the chain from the browser because the public RPCs allow it."
)
title(s, "Why BNB Chain")
box(s, 0.7, 1.6, 5.9, 0.9, "≈ 0.0000034 tBNB", 40, FG, True)
box(s, 0.7, 2.5, 5.9, 1.2, "per recorded decision: `33,189–34,323` gas\nat `0.1` gwei (our receipts, BSC Testnet)", 18, MUTED)
box(s, 0.7, 4.0, 5.9, 0.9, "≈ 0.00008 tBNB a day", 30, FG, True)
box(s, 0.7, 4.8, 5.9, 1.0, "for an hourly bot (24 decisions);\ndeploy: `449,207` gas ≈ 0.000045 tBNB", 18, MUTED)
box(s, 7.0, 1.6, 5.6, 0.6, "What else it gives us", 24, FG, True)
box(
    s, 7.0, 2.4, 5.6, 3.6,
    "EVM: Foundry tests and Sourcify verification on chain `97`\n"
    "Public RPCs a browser can call: the panel needs no backend\n"
    "Five smoke-test transactions within 132 blocks\n"
    "Mainnet (chain `56`) is a config switch; not deployed there",
    18,
)
box(s, 0.7, 6.5, 11.9, 0.5, "Testnet gas price, not a mainnet quote. Mainnet is on the roadmap, after an audit.", 14, MUTED)

# 8. Evidence and limits
s = slide("Numbers from real runs on 3 Oct 2026. CI is set up and runs after the next push; not claimed green yet.")
title(s, "Evidence and current limits")
facts = [
    ("30 / 30", "contract tests: unit, fuzz, invariant\n100% coverage of MacroGuard.sol", GREEN),
    ("48", "offline engine tests;\nno look-ahead tested as a property", GREEN),
    ("15", "web tests: contract reads and the\nwhat-if encoder against cast", GREEN),
    ("exact match", "source verified on Sourcify\n6 receipts, all status 1", GOLD),
]
for i, (big, small, color) in enumerate(facts):
    x = 0.7 + i * 3.1
    box(s, x, 1.6, 3.0, 0.8, big, 34, color, True)
    box(s, x, 2.5, 3.0, 1.4, small, 14, FG)
box(s, 0.7, 4.4, 11.9, 0.6, "Limits, said plainly", 22, AMBER, True)
box(
    s, 0.7, 5.0, 11.9, 1.9,
    "The contract never executes orders · the runner fails open if the RPC is down · the agent can `resume()` at once\n"
    "The drawdown is reported by the agent · no live bot tick and no profit claimed; backtests are research",
    16,
)

# 9. What Dex added
s = slide(PROVENANCE)
title(s, "What Dex added")
box(s, 0.7, 1.5, 5.2, 0.7, "Upstream DRIFT project", 24, MUTED, True)
box(s, 0.7, 2.3, 5.5, 2.8, "Quant engine, cockpit and MacroGuard.sol\n7 contract tests\nBuilt for a Mantle hackathon track (June 2026)\nMigrated to BNB Chain in bcc-ukdw/seed-bnb\n(commit `52671ce`, 29 Sep 2026)", 18)
box(s, 6.6, 1.5, 6.0, 0.7, "Dex's work in this fork", 24, FG, True)
box(
    s, 6.6, 2.3, 6.0, 4.0,
    "Own BSC Testnet deploy + 5-tx smoke test (status 1)\n"
    "Source verified on Sourcify (exact match)\n"
    "Public panel: live browser reads, receipts, limits\n"
    "Ask the contract: `eth_call` what-if, nothing signed\n"
    "+23 contract tests, 48 engine tests, 15 web tests, CI\n"
    "Labelled Binance data fallback · threat model\n"
    "Landing, visual system, PRD, this deck",
    17,
)
box(s, 0.7, 6.4, 11.9, 0.5, f"Code: `{REPO_URL.removeprefix('https://')}`", 16, MUTED)

# 10. Roadmap and links
s = slide("Roadmap items are not built yet. The links are live today; the BscScan and Sourcify links are clickable.")
title(s, "Roadmap")
box(s, 0.7, 1.3, 11.9, 0.5, "Not built yet; listed in order of trust gained per step", 15, AMBER, True)
box(
    s, 0.7, 1.9, 6.0, 2.2,
    "Resume through a multisig or timelock\nFail-closed mode for the runner\nSigned regime verdicts (EIP-712) with an input hash",
    18,
)
box(
    s, 6.9, 1.9, 6.0, 2.2,
    "Attested equity, so the drawdown is not self-reported\nRecord the raw signal plus a veto flag\nAgent identity (ERC-8004), then mainnet after an audit",
    18,
)
box(s, 0.7, 4.3, 11.9, 0.6, "Check it yourself", 24, FG, True)
link_row(s, 5.0, "Live demo", LIVE_URL, FG)
link_row(s, 5.5, "Code", REPO_URL, FG)
link_row(s, 6.0, "Contract", BSCSCAN_URL, GOLD, clickable=True)
link_row(s, 6.5, "Source", SOURCIFY_URL, GOLD, clickable=True)

# Document properties: the library's template ships its own author and description.
props = prs.core_properties
props.title = "DRIFT — a trading bot whose risk rules you can verify"
props.author = "Dex Bennett"
props.last_modified_by = "Dex Bennett"
props.comments = "DRIFT pitch deck: public risk gate on BNB Chain"  # dc:description
props.subject = "Indonesia Web3 Hackathon 2026 · BNB Chain"
props.created = props.modified = datetime.now(timezone.utc).replace(tzinfo=None)
props.revision = 4

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)

# The template's extended properties (docProps/app.xml) describe another file: a 4:3 deck
# with zero slides, saved by another application. Replace them with this deck's facts.
SLIDES = len(prs.slides._sldIdLst)
APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
    f"<PresentationFormat>Widescreen</PresentationFormat><Slides>{SLIDES}</Slides><Notes>{SLIDES}</Notes>"
    "<HiddenSlides>0</HiddenSlides><ScaleCrop>false</ScaleCrop><LinksUpToDate>false</LinksUpToDate>"
    "<SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged></Properties>"
)
# PowerPoint draws hyperlinks in the theme's link colour (template blue, unreadable on ink),
# not in the run colour; the only links in this deck are the on-chain ones, so the theme
# link colour becomes gold.
def patch(name: str, data: bytes) -> bytes:
    if name == "docProps/app.xml":
        return APP_XML.encode("utf-8")
    if name.startswith("ppt/theme/") and name.endswith(".xml"):
        text = data.decode("utf-8")
        text = re.sub(r"<a:hlink>.*?</a:hlink>", '<a:hlink><a:srgbClr val="F0B90B"/></a:hlink>', text)
        text = re.sub(r"<a:folHlink>.*?</a:folHlink>", '<a:folHlink><a:srgbClr val="F0B90B"/></a:folHlink>', text)
        return text.encode("utf-8")
    return data


tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
    for item in src.infolist():
        dst.writestr(item, patch(item.filename, src.read(item.filename)))
os.replace(tmp, OUT)

print("saved:", OUT)
print("bytes:", os.path.getsize(OUT))
print("slides:", SLIDES)
